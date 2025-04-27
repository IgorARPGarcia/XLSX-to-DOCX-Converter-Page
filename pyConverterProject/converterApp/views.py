from django.shortcuts import render
import tempfile
from django.http import HttpResponse
from django.views.decorators.csrf import csrf_exempt
import openpyxl
from docx import Document

from django.shortcuts import render

def index(request):
    return render(request, 'converterApp/index.html')

@csrf_exempt
def convert_xlsx_to_docx(request):
    if request.method == "POST":
        uploaded_file = request.FILES.get("file")
        if not uploaded_file:
            return HttpResponse("No file uploaded", status=400)

        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp_xlsx:
            for chunk in uploaded_file.chunks():
                temp_xlsx.write(chunk)
            temp_xlsx_path = temp_xlsx.name
            
            doc = Document()

            wb = openpyxl.load_workbook(temp_xlsx_path)
            sheet = wb.active

            for row in sheet.iter_rows(values_only=True):
                line = "\t".join([str(cell) if cell is not None else "" for cell in row])
                doc.add_paragraph(line)

            with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as temp_docx:
                doc.save(temp_docx.name)
                temp_docx.seek(0)
                response = HttpResponse(temp_docx.read(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document") 
                response["Content-Disposition"] = 'attachment; filename="converted.docx"'
                return response
    return HttpResponse("Only POST method allowed", status=405)
